import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell margins (padding) in dxa."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def add_styled_table(doc, headers, data, col_widths=None):
    """Creates a beautifully formatted table with teal header."""
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "0F766E")  # Teal primary
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
            run.font.name = "Segoe UI"

    # Format Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=150, right=150)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.name = "Segoe UI"
                run.font.color.rgb = RGBColor(51, 65, 85)

    # Set column widths if provided
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

    doc.add_paragraph()  # Spacing
    return table

def create_callout_box(doc, title, text, bg_hex="F0FDFA", border_hex="0D9488"):
    """Creates an alert/callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    # Left border only
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"★ {title}\n")
    run_t.bold = True
    run_t.font.name = "Segoe UI"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(15, 118, 110)

    run_body = p.add_run(text)
    run_body.font.name = "Segoe UI"
    run_body.font.size = Pt(9)
    run_body.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph()

def build_documentation():
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    r_title = title_p.add_run("HR NEXUS")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(26)
    r_title.font.color.rgb = RGBColor(15, 118, 110)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    r_sub = sub_p.add_run("Enterprise Human Resources, Attendance, Time-Off & Deterministic Payroll ERP System\nComprehensive Technical Architecture & Operations Specification")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata Ribbon
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(18)
    r_meta = meta_p.add_run("Version 2.0.0  |  Authorized Roles: Admin, HR Manager, Employee  |  Stack: React 19 + Express + MySQL 8.0")
    r_meta.bold = True
    r_meta.font.name = "Segoe UI"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(13, 148, 136)

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 1: EXECUTIVE SUMMARY & SYSTEM OVERVIEW
    # ──────────────────────────────────────────────────────────────────────────
    h1 = doc.add_heading("1. Executive Summary & Project Overview", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    for run in h1.runs:
        run.font.name = "Segoe UI"
        run.font.color.rgb = RGBColor(15, 118, 110)

    p1 = doc.add_paragraph(
        "HR Nexus is an enterprise-grade Human Resource Management (HRMS) and deterministic payroll execution platform. "
        "Engineered for accuracy, reliability, and clean corporate workflows, HR Nexus provides unified lifecycle governance "
        "covering employee records, structured contracts, working schedule configurations, real-time punch clock attendance, "
        "entitlement-tracked time-off requests, and high-precision multi-stage payroll cycles with PDF payslip generation."
    )
    p1.paragraph_format.space_after = Pt(8)

    create_callout_box(
        doc,
        "100% Dynamic Architecture Guarantee",
        "HR Nexus strictly enforces pure dynamic data sourcing. Zero mock records, zero static placeholders, and zero demo "
        "shims are permitted in production code. Every workforce metric, attendance percentage, leave balance, contract wage, "
        "and payroll calculation originates from and persists directly to the MySQL database via typed RESTful API endpoints."
    )

    doc.add_heading("Core Technology Stack", level=2)
    add_styled_table(
        doc,
        ["Layer", "Technology", "Architectural Role"],
        [
            ["Frontend Client", "React 19, TypeScript, Vite", "Dual-panel responsive layout, strict RBAC route guards, live punch clocks, and reactive dashboards."],
            ["Design & Styling", "Vanilla CSS3 Design System", "Curated Teal/Emerald corporate tokens (#0f766e, #0d9488), dark/light theme switching, and glassmorphism."],
            ["API Service", "Node.js, Express, TypeScript", "Centralized connection pooling, JWT session handling, parameter sanitization, and deterministic business logic."],
            ["Database Engine", "MySQL 8.0 (InnoDB)", "3NF/BCNF relational normalization, foreign key constraints, atomic transactions, and parameterized SQL queries."],
            ["Document Engine", "PDFKit", "Cryptographically stable, vector-rendered, immutable employee payslip PDF export."]
        ],
        col_widths=[Inches(1.5), Inches(2.2), Inches(2.8)]
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 2: THE 3 AUTHORIZED ROLES (ENTERPRISE RBAC)
    # ──────────────────────────────────────────────────────────────────────────
    h2 = doc.add_heading("2. Role-Based Access Control (RBAC) — Exactly 3 Authorized Roles", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    for run in h2.runs:
        run.font.name = "Segoe UI"
        run.font.color.rgb = RGBColor(15, 118, 110)

    p2 = doc.add_paragraph(
        "To ensure uncompromising data governance, audit integrity, and regulatory segregation of duties, HR Nexus "
        "is strictly governed by exactly three (3) authorization roles. Legacy granular roles have been unified into this "
        "clean, industry-standard corporate taxonomy:"
    )
    p2.paragraph_format.space_after = Pt(8)

    doc.add_heading("Role Descriptions & Authority Matrix", level=2)

    doc.add_paragraph(
        "1. Admin (Administrator / Superadmin):\n"
        "   • Executive system oversight, enterprise governance, and complete operational command.\n"
        "   • Unrestricted access across Workforce Directory, Contracts, Working Schedules, Salary Structures & Rules, System Audit, and Settings.\n"
        "   • Authorized to launch, compute, validate, and execute Payruns, inspect and download all company payslips.\n"
        "   • Controls system-level diagnostics, formula configuration rules, and platform parameters."
    )

    doc.add_paragraph(
        "2. HR Manager (Human Resources Lead):\n"
        "   • Operational personnel management and workforce lifecycle administration.\n"
        "   • Full management of the Employee Directory (adding personnel, modifying profiles, managing 360 Hubs, tracking active/probation/terminated statuses).\n"
        "   • Drafting and maintaining Employment Contracts and assigning Working Schedules.\n"
        "   • Supervision of company-wide Attendance records and shift compliance.\n"
        "   • Adjudication of Time-Off & Leave requests (approving or refusing requests with real-time balance reconciliation).\n"
        "   • Execution of periodic Payruns (Draft -> Computed -> Validated -> Paid) and verification of staff payslips."
    )

    doc.add_paragraph(
        "3. Employee (Staff Member / Self-Service):\n"
        "   • Employee Self-Service (ESS) workspace focused on personal workforce engagement.\n"
        "   • Live Interactive Punch Clock: daily limit of 1 check-in and 1 check-out per calendar day, real-time shift timer.\n"
        "   • Leave Balances & Time-Off Tracker: live visibility into Paid Time Off (PTO), Sick Leave, and Casual/Unpaid entitlements, with modal leave request submission.\n"
        "   • Latest Payslip Snapshot: real-time deposit breakdown, tax/deductions summary, and personal historical payslip archive with immutable PDF download.\n"
        "   • Company bulletins, notices, and personal working schedule visibility.\n"
        "   • Strict Security & Data Isolation: Employees cannot view, query, or modify other personnel's contracts, salaries, attendance records, or payslips."
    )

    doc.add_heading("Functional Domain Permissions Matrix", level=2)
    add_styled_table(
        doc,
        ["Functional Domain", "Admin", "HR Manager", "Employee", "Security & Governance Rules"],
        [
            ["Authentication & Login", "Full", "Full", "Full", "JWT tokens (HS256), email/password login, 0 demo quick-fill buttons."],
            ["Dashboard Experience", "Executive Operations Hub", "HR Operations Hub", "Employee Self-Service", "Adaptive role view: superadmin metrics vs personal punch/leave/payslip."],
            ["Employee Directory", "Full Read/Write", "Full Read/Write", "Forbidden", "Active employees first; terminated personnel sorted to the end."],
            ["Employee 360 Hub", "Full Read/Write", "Full Read/Write", "Self Profile Only", "Personal contact, contract wage, bank accounts, and attendance ledger."],
            ["Contracts Management", "Full Read/Write", "Full Read/Write", "Read (Own Wage)", "Employment terms, salary structures, weekly working schedules."],
            ["Working Schedules", "Full Read/Write", "Full Read/Write", "Read (Own Shift)", "Weekly hours configuration (Standard 40h, Part-time 20h)."],
            ["Attendance Registry", "Full Read/Write", "Full Read/Write", "Self Punch Clock", "1 In / 1 Out daily limit; live timer; attendance rate recalculation."],
            ["Leave & Time-Off", "Review / Approve / Refuse", "Review / Approve / Refuse", "Request / View Own", "Calendar balance deduction; instant status sync; reason validation."],
            ["Payrun Processing", "Execute & Validate", "Execute & Validate", "Forbidden", "4-Stage pipeline: DRAFT -> COMPUTED -> VALIDATED -> PAID."],
            ["Payslip Access", "Company-wide & PDF", "Company-wide & PDF", "Own Payslips & PDF", "Immutable salary records; download PDF export; zero cross-access."],
            ["Salary Structures & Rules", "Full Read/Write", "Full Read/Write", "Forbidden", "Basic (60%), HRA (25%), Allowances (15%), Tax & PF rules."],
            ["System Settings & Diagnostics", "Full Access", "Forbidden", "Forbidden", "Restricted to Superadmin level only."]
        ],
        col_widths=[Inches(1.4), Inches(0.9), Inches(1.1), Inches(1.1), Inches(2.0)]
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 3: DATABASE SCHEMA & RELATIONAL ARCHITECTURE
    # ──────────────────────────────────────────────────────────────────────────
    h3 = doc.add_heading("3. Database Architecture & Relational Schema (MySQL 8.0)", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    for run in h3.runs:
        run.font.name = "Segoe UI"
        run.font.color.rgb = RGBColor(15, 118, 110)

    p3 = doc.add_paragraph(
        "The underlying database follows strict 3NF/BCNF relational normalization to eliminate data anomalies and ensure "
        "audit compliance. All relational entities utilize InnoDB storage engine with parameterized query isolation."
    )
    p3.paragraph_format.space_after = Pt(8)

    add_styled_table(
        doc,
        ["Table Name", "Primary Key", "Foreign Keys", "Description & Responsibilities"],
        [
            ["employees", "id (VARCHAR/UUID)", "None", "Identity, name, email, department, position, gender, dateOfBirth, status (ACTIVE/PROBATION/TERMINATED), bank account."],
            ["working_schedules", "id (VARCHAR)", "None", "Schedule definitions, name (e.g. Standard 40h Regular), weekly target hours."],
            ["salary_structures", "id (VARCHAR)", "None", "Base salary classification structures and applicable currency definitions."],
            ["salary_rules", "id (VARCHAR)", "structure_id", "Component calculation rules: Category (BASIC/ALLOWANCE/GROSS/DEDUCTION/NET), Type (FIXED/PERCENTAGE/FORMULA)."],
            ["contracts", "id (VARCHAR)", "employee_id, structure_id, schedule_id", "Contract terms, monthly wage, start/end dates, active status."],
            ["attendance_records", "id (VARCHAR)", "employee_id", "Daily attendance logs, date, check_in, check_out, worked_hours, status (PRESENT/LATE/ABSENT/OVERTIME)."],
            ["time_off_requests", "id (VARCHAR)", "employee_id", "Leave applications, leave_type (PTO/Sick/Unpaid), start/end dates, duration_days, reason, status."],
            ["payruns", "id (VARCHAR)", "None", "Payroll run headers, cycle period, total_gross, total_net, employee_count, status (DRAFT/COMPUTED/VALIDATED/PAID)."],
            ["payslips", "id (VARCHAR)", "payrun_id, employee_id", "Detailed employee payslip line items, basic, hra, allowances, gross, deductions, net salary."]
        ],
        col_widths=[Inches(1.4), Inches(1.1), Inches(1.5), Inches(2.5)]
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 4: REAL-TIME DYNAMIC WORKFLOWS
    # ──────────────────────────────────────────────────────────────────────────
    h4 = doc.add_heading("4. Real-Time Dynamic Workforce Workflows", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    for run in h4.runs:
        run.font.name = "Segoe UI"
        run.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_heading("4.1 Live Punch Clock & Dynamic Attendance Rate Engine", level=2)
    p_att = doc.add_paragraph(
        "Attendance in HR Nexus is responsive and mathematically dynamic. Attendance rate is not a static number—it is "
        "calculated directly by the database engine across all historical shift records for an employee:\n\n"
        "• Attendance Rate Formula: ROUND((SUM(CASE WHEN status IN ('PRESENT', 'OVERTIME') THEN 1.0 WHEN status = 'LATE' THEN 0.8 ELSE 0.0 END) * 100.0) / COUNT(*))\n"
        "• Dynamic Movement: When an employee records a PRESENT shift, their attendance percentage dynamically increases. "
        "When an ABSENT record is recorded, their attendance percentage dynamically decreases.\n"
        "• Daily Limit: The employee punch clock enforces a strict enterprise rule of 1 check-in and 1 check-out per calendar day, "
        "preventing double punches while displaying real-time session elapsed time."
    )
    p_att.paragraph_format.space_after = Pt(8)

    doc.add_heading("4.2 Employee Directory Sorting & Date of Birth (DOB) Standard", level=2)
    p_emp = doc.add_paragraph(
        "• Terminated Employees Placement: The Employee Directory automatically sorts all personnel so that ACTIVE and PROBATION "
        "staff members appear first in operational view, while all TERMINATED employees are consistently placed at the very end of the list.\n"
        "• Date of Birth (DOB): Every employee record stores a validated Date of Birth. The system enforces an adult labor threshold "
        "(age >= 18 years) on employee creation and formats the date consistently across directory tables and 360 Hub modals."
    )
    p_emp.paragraph_format.space_after = Pt(8)

    doc.add_heading("4.3 Deterministic Payroll Engine & PDF Generation", level=2)
    p_pay = doc.add_paragraph(
        "• 4-Stage Lifecycle: Payruns transition immutably through DRAFT -> COMPUTED -> VALIDATED -> PAID.\n"
        "• Dynamic Cycle Initialization: The Payrun creation wizard automatically detects the current calendar date and defaults "
        "to the active month and year (e.g., October 2026 Regular Cycle: 2026-10-01 to 2026-10-31).\n"
        "• Calculation Rules: Base Salary (60% of wage), HRA (25%), Allowances (15%), Gross Pay, Unpaid Leave Prorated Deductions, "
        "Statutory Income Tax (10%), and Provident Fund/Deductions (7%).\n"
        "• Immutable PDF Payslips: Once a payrun is validated, personnel can download immutable, professionally branded PDF payslips "
        "generated server-side using PDFKit."
    )
    p_pay.paragraph_format.space_after = Pt(8)

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 5: SCREEN DIRECTORY & ROUTE SPECIFICATION
    # ──────────────────────────────────────────────────────────────────────────
    h5 = doc.add_heading("5. Screen Directory & Route Specifications", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    for run in h5.runs:
        run.font.name = "Segoe UI"
        run.font.color.rgb = RGBColor(15, 118, 110)

    add_styled_table(
        doc,
        ["Route", "Page Title", "Permitted Roles", "Core Screen Capabilities"],
        [
            ["/login", "Enterprise Authentication", "Public / Unauthenticated", "Clean email + password form, password visibility toggle, capability panel, zero static shortcuts."],
            ["/dashboard", "Operations / Self-Service Hub", "Admin, HR Manager, Employee", "Adaptive: Management KPI analytics & lifecycle tracker OR Employee punch clock & leave balances."],
            ["/employees", "Employee Directory & 360 Hub", "Admin, HR Manager", "Directory table, search, department filter, terminated-at-end sort, DOB, live attendance %, 360 Hub."],
            ["/contracts", "Employment Contracts", "Admin, HR Manager", "Active and historical contracts, salary structure linkages, wage assignments."],
            ["/schedules", "Working Schedules", "Admin, HR Manager", "Weekly hours, daily schedule specifications, shift templates."],
            ["/attendance", "Attendance Registry", "Admin, HR Manager, Employee", "Real-time attendance logs, check-in/out records, worked hours, status filters."],
            ["/time-off", "Time-Off & Leave Governance", "Admin, HR Manager, Employee", "Request leave modal (today defaults), manager approve/refuse actions, entitlement ledger."],
            ["/payruns", "Payroll Command Center", "Admin, HR Manager", "Dynamic monthly cycle wizard, compute batch, validate run, record disbursements."],
            ["/payslips", "Payslip Archive & PDF Hub", "Admin, HR Manager, Employee", "Historical payslip search, itemized breakdown modal, instant PDF download."],
            ["/salary-rules", "Salary Structures & Rule Engine", "Admin, HR Manager", "Configurable salary rules, calculation types (Fixed, Percentage, Formula)."],
            ["/settings", "System Settings & Diagnostics", "Admin (Superadmin only)", "Platform health checks, database status, security preferences."]
        ],
        col_widths=[Inches(1.1), Inches(1.4), Inches(1.4), Inches(2.6)]
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 6: LOCAL SETUP & RUNTIME EXECUTION GUIDE
    # ──────────────────────────────────────────────────────────────────────────
    h6 = doc.add_heading("6. Local Development & Operational Setup", level=1)
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)
    for run in h6.runs:
        run.font.name = "Segoe UI"
        run.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_heading("Environment Configuration", level=2)
    add_styled_table(
        doc,
        ["Variable", "Layer", "Default Value", "Description"],
        [
            ["PORT", "Backend", "5000", "Express REST API listening port."],
            ["DB_HOST", "Backend", "localhost", "MySQL server hostname."],
            ["DB_PORT", "Backend", "3306", "MySQL server port."],
            ["DB_USER", "Backend", "root", "MySQL administrative username."],
            ["DB_PASSWORD", "Backend", "[Configured]", "MySQL authentication password."],
            ["DB_NAME", "Backend", "peoplepay360", "Target schema database name."],
            ["JWT_SECRET", "Backend", "[Secure Token]", "HMAC-SHA256 secret key for signing session tokens."],
            ["JWT_EXPIRES_IN", "Backend", "24h", "Session authentication validity duration."],
            ["VITE_API_BASE_URL", "Frontend", "http://localhost:5000", "Base URL for client-side API requests."]
        ],
        col_widths=[Inches(1.5), Inches(0.9), Inches(1.5), Inches(2.6)]
    )

    doc.add_heading("Standard Execution Commands", level=2)
    add_styled_table(
        doc,
        ["Working Directory", "Command", "Operational Function"],
        [
            ["Root", "npm install", "Installs root-level packages and shared scripts."],
            ["server/", "npm run dev", "Launches backend API in hot-reload mode via tsx watch (Port 5000)."],
            ["server/", "npm run build", "Compiles TypeScript server code to dist/ via tsc."],
            ["server/", "npm run db:init", "Initializes MySQL tables and applies normalization schema."],
            ["client/", "npm run dev", "Starts Vite reactive development server (Port 5173)."],
            ["client/", "npm run build", "Generates optimized, minified production assets."]
        ],
        col_widths=[Inches(1.3), Inches(1.6), Inches(3.6)]
    )

    # Save Document
    filename = "HR_Nexus_Project_Documentation.docx"
    doc.save(filename)
    # Also save as legacy filename for compatibility
    doc.save("PeoplePay360_Project_Documentation.docx")
    print(f"[SUCCESS] Generated complete, updated documentation: {filename} and PeoplePay360_Project_Documentation.docx")

if __name__ == "__main__":
    build_documentation()
